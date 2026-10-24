
from io import BytesIO
import json
import re
import zipfile

import pymupdf
import pytesseract
import streamlit as st
from docx import Document
from PIL import Image
from pptx import Presentation


# Tesseract OCR path for Windows
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ZIP_FILES = 100
MAX_ZIP_MEMBER_BYTES = 10 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 50 * 1024 * 1024

SUPPORTED_DOCUMENTS = (".pdf", ".txt", ".md", ".docx", ".pptx")


def extract_pdf_text(file_bytes: bytes) -> tuple[str, int]:
    """Extract PDF text and OCR pages that have no selectable text."""
    parts = []
    ocr_pages = 0

    with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            page_text = page.get_text().strip()

            if page_text:
                parts.append(page_text)
                continue

            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(2, 2),
                alpha=False,
            )
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            page_text = pytesseract.image_to_string(image).strip()
            ocr_pages += 1

            if page_text:
                parts.append(page_text)

    return "\n\n".join(parts), ocr_pages


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and tables from DOCX files."""
    document = Document(BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def extract_pptx_text(file_bytes: bytes) -> str:
    """Extract text from PowerPoint slides and tables."""
    presentation = Presentation(BytesIO(file_bytes))
    parts = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_parts = [f"Slide {slide_number}"]

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                slide_parts.append(shape.text)

            if shape.has_table:
                for row in shape.table.rows:
                    slide_parts.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

        parts.append("\n".join(slide_parts))

    return "\n\n".join(parts)


def extract_text_by_name(
    filename: str,
    file_bytes: bytes,
) -> tuple[str, str]:
    """Extract text according to a supported file extension."""
    name = filename.lower()

    if name.endswith(".pdf"):
        text, ocr_pages = extract_pdf_text(file_bytes)
        return text, f"PDF; OCR attempted on {ocr_pages} page(s)"

    if name.endswith((".txt", ".md")):
        return (
            file_bytes.decode("utf-8-sig", errors="replace"),
            "Plain text",
        )

    if name.endswith(".docx"):
        return extract_docx_text(file_bytes), "Word document"

    if name.endswith(".pptx"):
        return extract_pptx_text(file_bytes), "PowerPoint presentation"

    raise ValueError(f"Unsupported document format: {filename}")


def extract_zip_documents(file_bytes: bytes) -> tuple[str, list[str]]:
    """Read supported documents from a ZIP without executing its contents."""
    parts = []
    included_files = []
    total_size = 0

    with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
        members = [
            item
            for item in archive.infolist()
            if not item.is_dir()
            and not item.filename.startswith("__MACOSX/")
        ]

        if len(members) > MAX_ZIP_FILES:
            raise ValueError(
                f"ZIP contains more than {MAX_ZIP_FILES} files."
            )

        for member in members:
            name = member.filename
            normalized_name = name.replace("\\", "/")
            path_parts = normalized_name.split("/")

            if (
                normalized_name.startswith("/")
                or ".." in path_parts
                or ":" in path_parts[0]
            ):
                continue

            if not name.lower().endswith(SUPPORTED_DOCUMENTS):
                continue

            if member.file_size > MAX_ZIP_MEMBER_BYTES:
                continue

            total_size += member.file_size

            if total_size > MAX_ZIP_TOTAL_BYTES:
                raise ValueError(
                    "The ZIP's supported documents exceed the size limit."
                )

            content = archive.read(member)
            extracted, _ = extract_text_by_name(name, content)

            if extracted.strip():
                parts.append(f"DOCUMENT: {name}\n{extracted}")
                included_files.append(name)

    return "\n\n".join(parts), included_files


def extract_uploaded_file(uploaded_file) -> tuple[str, str]:
    """Extract text from one uploaded document or ZIP archive."""
    file_bytes = uploaded_file.getvalue()

    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("The uploaded file exceeds the 25 MB limit.")

    if uploaded_file.name.lower().endswith(".zip"):
        text, filenames = extract_zip_documents(file_bytes)
        details = (
            "ZIP archive; documents read: "
            + (", ".join(filenames) if filenames else "none")
        )
        return text, details

    return extract_text_by_name(uploaded_file.name, file_bytes)


def review_project_document(text: str) -> dict:
    """Create a preliminary heuristic-based document audit."""
    cleaned = text.strip()

    if not cleaned:
        return {
            "status": "Insufficient content",
            "summary": "No readable text was extracted.",
            "findings": [
                {
                    "area": "Document content",
                    "status": "Needs attention",
                    "details": (
                        "No readable text was found. Check the file or "
                        "try a clearer document."
                    ),
                }
            ],
            "candidate_requirements": [],
            "limitations": [
                "No source code was inspected.",
                "This version uses keyword and pattern checks, not an AI model.",
            ],
            "extracted_text_preview": "",
        }

    lower_text = cleaned.lower()

    categories = {
        "Project objectives": [
            "objective",
            "goal",
            "purpose",
            "problem statement",
            "aim",
            "intended to",
        ],
        "Requirements": [
            "requirement",
            "must ",
            "shall ",
            "should ",
            "needs to",
            "has to",
        ],
        "Deliverables": [
            "deliverable",
            "submission",
            "expected output",
            "final product",
            "hand over",
        ],
        "Timeline": [
            "deadline",
            "timeline",
            "milestone",
            "schedule",
            "due date",
            "week ",
            "month ",
        ],
        "Risks and constraints": [
            "risk",
            "limitation",
            "constraint",
            "dependency",
            "assumption",
            "security concern",
        ],
        "Testing and acceptance": [
            "test plan",
            "testing",
            "validation",
            "acceptance criteria",
            "success criteria",
            "expected result",
            "pass criteria",
        ],
    }

    findings = []

    for area, keywords in categories.items():
        matches = [
            term.strip()
            for term in keywords
            if term in lower_text
        ]

        if matches:
            findings.append(
                {
                    "area": area,
                    "status": "Related content found",
                    "details": (
                        "Matched wording: "
                        + ", ".join(matches)
                        + ". Check the surrounding text to confirm "
                        "the information is sufficient."
                    ),
                }
            )
        else:
            findings.append(
                {
                    "area": area,
                    "status": "Review recommended",
                    "details": (
                        f"No obvious wording about {area.lower()} was found. "
                        "It may be missing or described differently."
                    ),
                }
            )

    # Extract sentences that might describe requirements.
    candidates = []
    seen = set()

    for line in re.split(r"[\r\n]+", cleaned):
        line = re.sub(
            r"^\s*(?:[-*•]+|\d+[.)])\s*",
            "",
            line,
        ).strip()

        if not line:
            continue

        sentences = re.split(r"(?<=[.!?])\s+", line)

        for sentence in sentences:
            sentence = sentence.strip()

            if len(sentence.split()) < 4 or len(sentence) > 400:
                continue

            sentence_lower = sentence.lower()

            looks_like_requirement = bool(
                re.search(
                    r"\b(must|shall|should|needs to|has to|required to)\b",
                    sentence_lower,
                )
                or re.search(
                    r"\b(the system|the application|users?|admins?)\b",
                    sentence_lower,
                )
                or re.search(
                    r"^\s*(create|allow|provide|support|enable|store|"
                    r"display|generate|prevent|ensure|manage|track|"
                    r"send|upload|download)\b",
                    sentence_lower,
                )
            )

            if looks_like_requirement:
                key = sentence_lower.rstrip(".!?")

                if key not in seen:
                    seen.add(key)
                    candidates.append(sentence)

    vague_terms = [
        "user-friendly",
        "fast",
        "easy to use",
        "secure",
        "robust",
        "efficient",
        "scalable",
        "good performance",
        "intuitive",
        "reliable",
    ]

    vague_matches = [
        term for term in vague_terms if term in lower_text
    ]

    if vague_matches:
        findings.append(
            {
                "area": "Potentially vague wording",
                "status": "Review recommended",
                "details": (
                    "These terms may need measurable definitions: "
                    + ", ".join(vague_matches)
                    + ". Consider defining targets or test conditions."
                ),
            }
        )

    if not candidates:
        findings.append(
            {
                "area": "Requirement clarity",
                "status": "Review recommended",
                "details": (
                    "No candidate requirements were extracted automatically. "
                    "Check whether the document states specific behaviours."
                ),
            }
        )

    word_count = len(cleaned.split())
    line_count = sum(
        1 for line in cleaned.splitlines() if line.strip()
    )

    return {
        "status": "Preliminary review complete",
        "summary": (
            f"Reviewed approximately {word_count} words across "
            f"{line_count} non-empty lines. Extracted "
            f"{len(candidates)} candidate requirement(s)."
        ),
        "findings": findings,
        "candidate_requirements": candidates,
        "limitations": [
            "Analysis uses keywords and patterns, not semantic AI reasoning.",
            "Missing keywords do not prove that information is absent.",
            "Candidate requirements may include descriptive sentences.",
            "The app does not inspect code or verify that features work.",
            "Embedded images in DOCX and PPTX are not OCR-processed.",
        ],
        "extracted_text_preview": cleaned[:5000],
    }


def main():
    st.set_page_config(
        page_title="ProofSprint AI | Document Auditor",
        page_icon="🔎",
        layout="wide",
    )

    st.title("ProofSprint AI")
    st.subheader("Project Document Auditor")

    st.write(
        "Upload one project document to review its objectives, requirements, "
        "deliverables, timeline, risks, and testing information."
    )

    st.info(
        "Document-only mode reviews written information. "
        "It does not verify whether project features work."
    )

    uploaded_file = st.file_uploader(
        "Upload one file",
        type=["pdf", "pptx", "docx", "txt", "md", "zip"],
        accept_multiple_files=False,
    )

    if uploaded_file is None:
        st.caption(
            "Supported formats: PDF, PPTX, DOCX, TXT, MD, and ZIP."
        )
        return

    st.write(f"**Selected file:** {uploaded_file.name}")

    if st.button("Run document audit", type="primary"):
        try:
            with st.spinner("Extracting and reviewing document..."):
                text, extraction_details = extract_uploaded_file(
                    uploaded_file
                )

                if not text.strip():
                    st.session_state.pop("document_audit_report", None)
                    st.warning(
                        "No readable document text was found. For ZIP files, "
                        "include supported PDF, DOCX, PPTX, TXT, or MD files."
                    )
                    return

                report = review_project_document(text)
                report["filename"] = uploaded_file.name
                report["extraction_details"] = extraction_details

                st.session_state["document_audit_report"] = report

        except (
            ValueError,
            zipfile.BadZipFile,
            RuntimeError,
            OSError,
        ) as exc:
            st.session_state.pop("document_audit_report", None)
            st.error(f"Could not process the file: {exc}")
            return

        except Exception as exc:
            st.session_state.pop("document_audit_report", None)
            st.error(
                f"An unexpected error occurred: {exc}. "
                "Check the file and try again."
            )
            return

    report = st.session_state.get("document_audit_report")

    if report and report.get("filename") == uploaded_file.name:
        st.divider()
        st.header("Audit report")
        st.write(f"**Status:** {report['status']}")
        st.write(report["summary"])
        st.write(f"**Extraction:** {report['extraction_details']}")

        st.subheader("Findings")

        for finding in report["findings"]:
            st.markdown(
                f"**{finding['area']}** — {finding['status']}"
            )
            st.write(finding["details"])

        st.subheader("Candidate requirements")

        requirements = report.get("candidate_requirements", [])

        if requirements:
            for number, requirement in enumerate(requirements, start=1):
                st.write(f"{number}. {requirement}")
        else:
            st.info(
                "No candidate requirements were extracted. "
                "Review the document manually."
            )

        st.subheader("Limitations")

        for limitation in report["limitations"]:
            st.write(f"- {limitation}")

        with st.expander("View extracted text preview"):
            st.text(report.get("extracted_text_preview", ""))

        st.download_button(
            "Download report as JSON",
            data=json.dumps(
                report,
                indent=2,
                ensure_ascii=False,
            ),
            file_name="proofsprint_document_audit.json",
            mime="application/json",
        )


if __name__ == "__main__":
    main()
